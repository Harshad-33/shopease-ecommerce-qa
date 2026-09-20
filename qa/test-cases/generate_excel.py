import csv
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_excel_test_cases():
    csv_file = os.path.join(os.path.dirname(__file__), 'test_cases.csv')
    xlsx_file = os.path.join(os.path.dirname(__file__), 'ShopEase_Manual_Test_Cases.xlsx')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Manual Test Cases"

    # Header styling
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Arial", size=10)
    pass_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
    pass_font = Font(name="Arial", size=10, bold=True, color="0F5132")
    high_font = Font(name="Arial", size=10, bold=True, color="842029")
    
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    with open(csv_file, mode='r', encoding='utf-8') as f:
        reader = csv.reader(f)
        for r_idx, row in enumerate(reader, 1):
            ws.append(row)
            for c_idx in range(1, len(row) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                
                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cell.font = cell_font
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    # Status column styling (Column 9)
                    if c_idx == 9 and cell.value == "PASS":
                        cell.fill = pass_fill
                        cell.font = pass_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    # Priority column styling (Column 10)
                    elif c_idx == 10:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        if cell.value == "High":
                            cell.font = high_font

    # Set column widths
    column_widths = {
        'A': 15,  # Test Case ID
        'B': 18,  # Module
        'C': 32,  # Test Scenario
        'D': 28,  # Preconditions
        'E': 38,  # Test Steps
        'F': 25,  # Test Data
        'G': 35,  # Expected Result
        'H': 32,  # Actual Result
        'I': 12,  # Status
        'J': 12   # Priority
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'

    wb.save(xlsx_file)
    print(f"Successfully generated formatted Excel test cases: {xlsx_file}")

if __name__ == "__main__":
    create_excel_test_cases()
